"""RoleFit Auditor v1.0: transparent, offline résumé/job-post textual alignment.

The generated score is a reproducible heuristic, NOT a score from an employer's ATS,
a hiring probability, an assessment of personal worth, or a validated selection test.
"""
from __future__ import annotations

from dataclasses import dataclass, asdict
from io import BytesIO
from pathlib import Path
import re
import unicodedata
import zipfile
from collections import Counter
from typing import Optional

CATEGORY_WEIGHTS = {'Required': 55.0, 'Preferred': 20.0, 'Responsibilities': 25.0}
VERSION = '1.0.0'

# Deliberately conservative: close textual equivalents, not inferred skill equivalence.
ALIASES = {
    'excel': ['microsoft excel', 'ms excel', 'excel'],
    'powerpoint': ['microsoft powerpoint', 'ms powerpoint', 'powerpoint'],
    'word': ['microsoft word', 'ms word'],
    'power bi': ['powerbi', 'power bi'],
    'sql': ['structured query language', 'sql'],
    'python': ['python'],
    'r programming': ['r programming', 'r studio', 'rstudio', 'r language'],
    'salesforce': ['salesforce'],
    'crm': ['customer relationship management', 'crm'],
    'healthcare': ['health care', 'healthcare'],
    'go-to-market': ['go to market', 'go-to-market', 'gtm'],
    'stakeholder management': ['stakeholder management', 'stakeholder engagement'],
    'cross-functional': ['cross functional', 'cross-functional'],
    'account management': ['account management', 'client portfolio management'],
    'client relationship management': ['client relationship management', 'customer relationship management'],
    'consultative selling': ['consultative selling', 'solution selling'],
    'business development': ['business development'],
    'project management': ['project management'],
    'financial modelling': ['financial modeling', 'financial modelling'],
    'data analysis': ['data analysis', 'data analytics'],
    'market research': ['market research'],
    'qualitative research': ['qualitative research', 'qualitative analysis'],
    'quantitative research': ['quantitative research', 'quantitative analysis'],
    'machine learning': ['machine learning'],
    'artificial intelligence': ['artificial intelligence', 'ai'],
    'bachelors degree': ["bachelor's degree", 'bachelors degree', 'bachelor of', 'bsc', 'b.a.', 'b.s.', 'bhsc'],
    'masters degree': ["master's degree", 'masters degree', 'master of', 'mba', 'msc', 'm.sc.'],
    'mba': ['mba', 'master of business administration'],
    'phd': ['phd', 'ph.d.', 'doctor of philosophy'],
    'agile': ['agile'],
    'scrum': ['scrum'],
    'tableau': ['tableau'],
}
# Avoid overinterpreting common words, including vague "strong communication" clauses.
STOP = set('''a an the and or of to in on with for from at as by is are be being been
we you your our their this that those these which who where about plus including
have has having will should can able ability skills skill required requirement
requirements preferred desired ideally must minimum demonstrated proven strong excellent
good exceptional solid prior relevant significant experience experiences experienced
knowledge familiarity understanding proficient proficiency competency competencies
background work working worked role roles team teams company organization organisation
individual candidate candidates years year yr yrs level high highly successful
necessary ideal applicant responsible responsibility duties qualifications qualified
provide providing ensure ensuring support supporting use using used through across
complex multiple various different wide variety degree equivalent ability equivalent
interest interested comfortable capable demonstrated previous opportunity
'''.split())
ACTIVITY = set('''lead led manage managed develop developed build built deliver delivered
analyze analysed analyzed collaborate collaborated coordinate coordinated drive drove
design designed implement implemented oversee oversaw support supported advise advised
identify identified present presented report reported plan planned create created
execute executed own owned improve improved conduct conducted communicate communicated
research researched evaluate evaluated translate translated'''.split())
TECH_SKILLS = {'excel', 'powerpoint', 'word', 'power bi', 'sql', 'python',
               'r programming', 'salesforce', 'tableau'}

MANUAL_RE = re.compile(
    r'\b(?:\d+\s*\+?\s*(?:years?|yrs?)|work authori[sz]ation|visa|sponsorship|'
    r'legal right to work|eligible to work|security clearance|citizenship|'
    r'valid driving licen[cs]e|willing(?:ness)? to (?:travel|relocate)|'
    r'located in|based in|fluen(?:t|cy) in|native speaker)\b', re.I)

HEADINGS = {
    'Required': [
        'minimum qualifications', 'minimum requirements', 'required qualifications',
        'essential criteria', 'essential qualifications', 'essential skills',
        'requirements', 'qualifications', 'what you bring', 'what we are looking for',
        'what we re looking for', 'who you are', 'your experience', 'about you',
        'skills and experience', 'skills & experience', 'must have', 'your profile',
    ],
    'Preferred': [
        'preferred qualifications', 'preferred requirements', 'preferred skills',
        'nice to have', 'nice to haves', 'desirable', 'bonus points', 'advantageous',
        'additional qualifications', 'would be a plus', 'ideally you have',
    ],
    'Responsibilities': [
        'responsibilities', 'key responsibilities', 'what you will do', 'what youll do',
        'what you ll do', 'the role', 'your responsibilities', 'key duties',
        'about the role', 'your impact', 'day to day', 'what you do',
        'in this role', 'duties', 'the opportunity',
    ],
}
EXCLUDED_HEADINGS = (
    'about us', 'about the company', 'who we are', 'benefits', 'salary and benefits', 'our benefits',
    'what we offer', 'compensation', 'salary', 'pay and benefits', 'perks',
    'diversity', 'equal opportunity', 'eeo', 'how to apply', 'application process',
    'privacy', 'accommodations', 'our mission', 'culture', 'our values',
)
REQUIREMENT_HINTS = re.compile(
    r'\b(experience|skills|ability|knowledge|proficien|familiar|expertise|'
    r'qualification|degree|certification|track record|communication|'
    r'analys[ie]|stakeholder|lead|develop|deliver|manage|collaborat|'
    r'understanding|research|work with|design|implement|responsib|'
    r'customer|client|strategy|planning|financial|analytics|marketing)\b', re.I)


@dataclass
class ParsedDocument:
    text: str
    filename: str
    extension: str
    issues: list[str]
    stats: dict


@dataclass
class Criterion:
    requirement: str
    category: str
    include: bool = True
    importance: float = 1.0


def normalize(text: str) -> str:
    s = unicodedata.normalize('NFKD', str(text)).encode('ascii', 'ignore').decode('ascii').lower()
    s = s.replace('&', ' and ').replace('–', '-').replace('—', '-')
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9+.#/-]+', ' ', s)).strip()


def tokens(text: str) -> list[str]:
    return re.findall(r'[a-z][a-z0-9]*|\b\d+\b', normalize(text))


def contains_phrase(text: str, phrase: str) -> bool:
    src = normalize(text)
    term = normalize(phrase)
    if not term:
        return False
    # Word boundaries prevent "java" matching "javascript" and "R" matching random letters.
    return bool(re.search(r'(?<![a-z0-9])' + re.escape(term) + r'(?![a-z0-9])', src))


def extract_document(file_bytes: bytes, filename: str) -> ParsedDocument:
    ext = Path(filename).suffix.lower()
    issues = []
    stats = {'pages': None, 'tables': 0, 'images': 0, 'word_count': 0, 'characters': 0}
    if ext == '.pdf':
        import fitz
        try:
            doc = fitz.open(stream=file_bytes, filetype='pdf')
        except Exception as exc:
            raise ValueError('Could not open the PDF. It may be encrypted or damaged.') from exc
        stats['pages'] = len(doc)
        parts = []
        for page in doc:
            parts.append(page.get_text('text', sort=True))
            stats['images'] += len(page.get_images(full=True))
        text = '\n'.join(parts)
        if len(tokens(text)) < 35:
            issues.append('Very little extractable PDF text. This may be a scanned/image-only document; OCR is not performed.')
        if stats['images'] > 5 and len(tokens(text)) < 160:
            issues.append('Image-heavy PDF: visually important content may not have been extracted.')
        doc.close()
    elif ext == '.docx':
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        try:
            doc = Document(BytesIO(file_bytes))
        except Exception as exc:
            raise ValueError('Could not open the DOCX file. Convert old .doc files to .docx first.') from exc
        output = []
        for child in doc.element.body.iterchildren():
            if child.tag == qn('w:p'):
                para = Paragraph(child, doc)
                val = para.text.strip()
                if val:
                    if para.style and 'list' in para.style.name.lower() and not val.startswith(('•', '-', '*')):
                        val = '• ' + val
                    output.append(val)
            elif child.tag == qn('w:tbl'):
                table = Table(child, doc)
                stats['tables'] += 1
                for row in table.rows:
                    output.append(' | '.join(c.text.strip().replace('\n', ' / ') for c in row.cells if c.text.strip()))
        text = '\n'.join(output)
        if stats['tables']:
            issues.append(f'DOCX contains {stats["tables"]} table(s). Table contents were extracted, but some ATS products may read them differently.')
        if any(section.header.paragraphs and any(p.text.strip() for p in section.header.paragraphs) for section in doc.sections):
            issues.append('Text exists in a Word header; some parsers may omit header content.')
        if any(section.footer.paragraphs and any(p.text.strip() for p in section.footer.paragraphs) for section in doc.sections):
            issues.append('Text exists in a Word footer; some parsers may omit footer content.')
        try:
            with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
                word_xml = (name for name in archive.namelist() if name.startswith('word/') and name.endswith('.xml'))
                if any(b'txbxContent' in archive.read(name) for name in word_xml):
                    issues.append('Word text boxes detected: their contents may not be included in the extracted text.')
        except zipfile.BadZipFile:
            pass
    elif ext in ('.txt', '.md'):
        text = file_bytes.decode('utf-8-sig', errors='replace')
    else:
        raise ValueError(f'Unsupported file type {ext or "unknown"}. Use PDF, DOCX, TXT or MD.')
    text = text.replace('\x00', '').strip()
    stats['word_count'] = len(tokens(text))
    stats['characters'] = len(text)
    if not text:
        issues.append('No text could be extracted.')
    replacement_fraction = text.count('�') / max(len(text), 1)
    if replacement_fraction > .005:
        issues.append('Unrecognized replacement characters suggest encoding/extraction problems.')
    return ParsedDocument(text=text, filename=filename, extension=ext, issues=issues, stats=stats)


def parsing_diagnostic(parsed: ParsedDocument) -> tuple[int, list[str]]:
    """Document extraction check, NOT predicted compatibility with any ATS."""
    warnings = list(parsed.issues)
    score = 100
    n = parsed.stats['word_count']
    if n < 40:
        score -= 65
    elif n < 150:
        score -= 25
    elif n > 2800:
        score -= 10
        warnings.append('Unusually long resume: review the extracted text for unwanted duplication.')
    if parsed.extension == '.pdf' and parsed.stats['pages'] and parsed.stats['pages'] > 3:
        score -= 8
    if parsed.stats['tables']:
        score -= min(parsed.stats['tables'] * 5, 15)
    if any('header' in i.lower() or 'footer' in i.lower() for i in warnings):
        score -= 8
    if any('text box' in i.lower() for i in warnings):
        score -= 12
    if any('replacement characters' in i.lower() for i in warnings):
        score -= 20
    if any('image-only' in i.lower() for i in warnings):
        score -= 50
    headings = [x for x in ['experience', 'education', 'skills', 'projects'] if contains_phrase(parsed.text, x)]
    if n >= 150 and len(headings) == 0:
        score -= 8
        warnings.append('No standard section headers detected in extracted text; check the reading order manually.')
    return max(0, score), warnings


def heading_category(line: str) -> Optional[str]:
    norm = normalize(line).rstrip(':')
    if len(norm) > 55:
        return None
    for group, headings in HEADINGS.items():
        if norm in headings:
            return group
    return None


def looks_excluded(line: str) -> bool:
    norm = normalize(line)
    return any(norm == h or norm.startswith(h + ':') for h in EXCLUDED_HEADINGS)


def clean_bullet(line: str) -> str:
    return re.sub(r'^\s*(?:[•●▪◦*\-–—]|\d+[.)])\s*', '', str(line).strip()).strip()


def infer_category(text: str, current: str) -> str:
    s = normalize(text)
    if re.search(r'\b(preferred|nice to have|bonus|desirable|advantageous|ideally)\b', s):
        return 'Preferred'
    if re.search(r'\b(required|must have|essential|minimum|mandatory)\b', s):
        return 'Required'
    return current


def extract_requirements(job_text: str, max_items: int = 45) -> list[Criterion]:
    """Propose requirements; human must review before treating resulting score as informative."""
    lines = [x.strip() for x in str(job_text).splitlines() if x.strip()]
    found: list[Criterion] = []
    current = None
    excluded = False
    seen = set()
    for i, original in enumerate(lines):
        line = clean_bullet(original)
        if not line:
            continue
        hd = heading_category(line)
        if hd:
            current, excluded = hd, False
            continue
        if looks_excluded(line) and len(line) < 65:
            current, excluded = None, True
            continue
        if excluded:
            continue
        if len(line) < 14 or len(line) > 420:
            continue
        if re.search(r'(https?://|@\w+\.\w+|\bapply now\b|equal opportunity|salary range|pay range|benefits include)', line, re.I):
            continue
        bullet = bool(re.match(r'^\s*(?:[•●▪◦*\-–—]|\d+[.)])\s+', original))
        if current is None and not REQUIREMENT_HINTS.search(line):
            continue
        # Within a named section, accept prose when line carries a plausible criterion.
        if current and not bullet and len(line) > 230 and not REQUIREMENT_HINTS.search(line):
            continue
        category = infer_category(line, current or ('Responsibilities' if re.match(r'(?i)^(lead|manage|develop|build|analy[sz]e|deliver|collaborate|design|support|drive|implement|conduct|create|execute)', line) else 'Required'))
        normalized = normalize(line)
        if normalized in seen:
            continue
        seen.add(normalized)
        found.append(Criterion(requirement=line.rstrip(' .;'), category=category))
        if len(found) >= max_items:
            break
    if not found:
        # Fallback for full-paragraph job ads (still user-editable).
        sentences = re.split(r'(?<=[.!?])\s+', ' '.join(lines))
        for sentence in sentences:
            sentence = sentence.strip()
            if 20 <= len(sentence) <= 240 and REQUIREMENT_HINTS.search(sentence):
                found.append(Criterion(sentence, infer_category(sentence, 'Required')))
            if len(found) >= max_items:
                break
    return found


def make_resume_units(resume: str) -> list[str]:
    units = []
    for line in resume.splitlines():
        s = clean_bullet(line)
        if len(s) < 8:
            continue
        # Keep whole bullets to retain context, split only very long paragraphs.
        if len(s) > 420:
            units.extend(x.strip() for x in re.split(r'(?<=[.!?])\s+', s) if len(x.strip()) >= 8)
        else:
            units.append(s)
    return units


def skill_names(text: str) -> list[str]:
    return [key for key, forms in ALIASES.items() if any(contains_phrase(text, form) for form in forms)]


def key_terms(text: str) -> set[str]:
    words = tokens(text)
    # Light suffix normalization, explicitly not semantic equivalence.
    def stem(w):
        if len(w) >= 7 and w.endswith('ing'): return w[:-3]
        if len(w) >= 7 and w.endswith('ed'): return w[:-2]
        if len(w) >= 6 and w.endswith('ies'): return w[:-3] + 'y'
        if len(w) >= 6 and w.endswith('s') and not w.endswith('ss'): return w[:-1]
        return w
    return {stem(w) for w in words if (len(w) >= 3 or w.isnumeric()) and w not in STOP and w not in ACTIVITY}


def is_negated(unit: str, phrase: str) -> bool:
    norm = normalize(unit)
    for m in re.finditer(r'(?<![a-z0-9])' + re.escape(normalize(phrase)) + r'(?![a-z0-9])', norm):
        before = norm[max(0, m.start() - 42):m.start()]
        if re.search(r'\b(no|not|without|lack|lacking|never)\b(?:\W+\w+){0,4}\W*$', before):
            return True
    return False


def relevant_skills_in_unit(unit: str, expected: list[str]) -> list[str]:
    present = []
    for key in expected:
        forms = ALIASES[key]
        if any(contains_phrase(unit, alias) and not is_negated(unit, alias) for alias in forms):
            present.append(key)
    return present


def match_criterion(requirement: str, resume_units: list[str], strict: bool = False) -> dict:
    if MANUAL_RE.search(requirement):
        return {'status': 'Manual review', 'credit': None, 'evidence': '', 'reason': 'Eligibility, tenure or similar condition cannot be reliably verified from resume text.', 'terms': []}
    req_norm = normalize(requirement)
    req_keys = key_terms(requirement)
    expected_skills = skill_names(requirement)
    best = None
    best_score = -1.0
    for unit in resume_units:
        unit_norm = normalize(unit)
        unit_keys = key_terms(unit)
        shared = len(req_keys & unit_keys)
        coverage = shared / len(req_keys) if req_keys else 0
        present = relevant_skills_in_unit(unit, expected_skills)
        skill_coverage = len(present) / len(expected_skills) if expected_skills else None
        # Require a plausible skill or >=2 meaningful words; common vague wording is not evidence.
        has_minimum = bool(present) or shared >= 2 or (len(req_keys) == 1 and shared == 1)
        direct = bool(len(req_norm) >= 9 and contains_phrase(unit, requirement) and not is_negated(unit, requirement))
        if expected_skills and skill_coverage == 0:
            direct = False
        # A negated skill must not earn credit via generic word overlap.
        if (expected_skills and skill_coverage == 0) or re.match(r'^\s*(?:no|not|without|lacking)\b', unit_norm):
            credit, status, reason = 0.0, 'Not evidenced', 'Only absent or negated terms were found.'
        elif strict:
            if direct:
                credit, status, reason = 1.0, 'Direct', 'Full requirement phrase found in resume.'
            elif expected_skills and skill_coverage == 1 and coverage >= .34:
                credit, status, reason = .65, 'Related evidence', 'All specified technical terms found in one resume passage; scope still needs review.'
            elif expected_skills and skill_coverage > 0:
                credit, status, reason = .35, 'Partial', 'Some specified terms found in one resume passage.'
            else:
                credit, status, reason = 0.0, 'Not evidenced', 'No sufficiently specific textual evidence.'
        else:
            if direct:
                credit, status, reason = 1.0, 'Direct', 'Full requirement phrase found in resume.'
            elif has_minimum and expected_skills and skill_coverage == 1 and (coverage >= .50 or len(req_keys) <= 3):
                credit, status, reason = .85, 'Strong evidence', 'All named terms and relevant wording found in one resume passage.'
            elif has_minimum and coverage >= .70 and shared >= 2 and (skill_coverage is None or skill_coverage >= .75):
                credit, status, reason = .8, 'Strong evidence', 'Strong overlap of meaningful requirement terms in one resume passage.'
            elif has_minimum and (coverage >= .36 or (skill_coverage is not None and skill_coverage >= .5)):
                credit, status, reason = .45, 'Partial', 'Some aligned terms in resume; requirement may not be fully demonstrated.'
            elif has_minimum and expected_skills and skill_coverage:
                credit, status, reason = .3, 'Partial', 'Named term present but broader requirement is not demonstrated.'
            else:
                credit, status, reason = 0.0, 'Not evidenced', 'No sufficiently specific textual evidence.'
        # Prefer more complete coverage and relevant evidence, never keyword frequency.
        rank = credit + min(coverage, 1) * .001
        if rank > best_score:
            best_score = rank
            best = {'status': status, 'credit': credit,
                    'evidence': unit if credit else '', 'reason': reason,
                    'terms': sorted(set(present) | (req_keys & unit_keys)) if credit else []}
    # Multiple named software requirements may reasonably be evidenced across
    # separate bullets. Do not extend this aggregation to broader competencies.
    if len(expected_skills) > 1 and set(expected_skills).issubset(TECH_SKILLS):
        coverage_by_skill = {}
        for unit in resume_units:
            for skill in relevant_skills_in_unit(unit, expected_skills):
                coverage_by_skill.setdefault(skill, unit)
        if len(coverage_by_skill) == len(expected_skills) and (best is None or best['credit'] < .75):
            excerpts = list(dict.fromkeys(coverage_by_skill.values()))
            return {'status': 'Related evidence', 'credit': .75 if not strict else .65,
                    'evidence': ' | '.join(excerpts[:4]),
                    'reason': 'All named software/tools are mentioned, across separate resume passages. Proficiency still requires review.',
                    'terms': sorted(coverage_by_skill)}
    return best or {'status': 'Not evidenced', 'credit': 0.0, 'evidence': '', 'reason': 'No resume text.', 'terms': []}


def evaluate(criteria: list[dict], resume_text: str, strict: bool = False) -> dict:
    units = make_resume_units(resume_text)
    results = []
    for record in criteria:
        if not record.get('include', True):
            continue
        category = record.get('category', 'Required')
        if category not in CATEGORY_WEIGHTS:
            continue
        requirement = str(record.get('requirement', '')).strip()
        if len(requirement) < 3:
            continue
        importance = max(0.0, min(5.0, float(record.get('importance', 1.0))))
        if importance <= 0:
            continue
        result = match_criterion(requirement, units, strict)
        results.append({**result, 'requirement': requirement,
                        'category': category, 'importance': importance})
    section = {}
    total_active_weight = 0.0
    total_weighted_credit = 0.0
    for category, nominal in CATEGORY_WEIGHTS.items():
        auto = [r for r in results if r['category'] == category and r['credit'] is not None]
        manual = [r for r in results if r['category'] == category and r['credit'] is None]
        importance_sum = sum(r['importance'] for r in auto)
        mean = (sum(r['importance'] * r['credit'] for r in auto) / importance_sum) if importance_sum else None
        section[category] = {'score': round(mean * 100, 1) if mean is not None else None,
                             'assessed': len(auto), 'manual': len(manual), 'weight': nominal}
        if mean is not None:
            total_active_weight += nominal
            total_weighted_credit += nominal * mean
    score = round(100 * total_weighted_credit / total_active_weight) if total_active_weight else None
    return {'score': score, 'sections': section, 'results': results,
            'assessed': sum(r['credit'] is not None for r in results),
            'manual': sum(r['credit'] is None for r in results),
            'mode': 'Strict' if strict else 'Balanced',
            'version': VERSION}


def recommendations(evaluation: dict, limit: int = 6) -> list[dict]:
    missing = [r for r in evaluation['results'] if r['credit'] is not None and r['credit'] < .8]
    category_rank = {'Required': 0, 'Responsibilities': 1, 'Preferred': 2}
    missing.sort(key=lambda r: (category_rank[r['category']], -r['importance'], r['credit']))
    out = []
    for item in missing[:limit]:
        advice = ('If you have performed this work, show it explicitly in an accurate experience/project bullet; otherwise treat it as a genuine gap.'
                  if item['credit'] == 0 else
                  'Review the supporting bullet. Clarify the scope, tools, actions and outcome only if those details are true.')
        out.append({'requirement': item['requirement'], 'category': item['category'], 'advice': advice})
    return out
